"""Reference-based Word skeleton for the MIT software office SITOP.

Requires python-docx from the document runtime. Does not modify source templates.
Import new_document/add_orientation and add project content to the returned Document.
"""
from copy import deepcopy
from pathlib import Path
import argparse
from datetime import date
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def _clone_part_content(source,target):
    """Copy content and remap relationship ids, including embedded brand images."""
    for child in list(target._element):target._element.remove(child)
    ids={}
    for rel in source.part.rels.values():
        if rel.reltype==RT.IMAGE and not rel.is_external:
            ids[rel.rId],_=target.part.get_or_add_image(BytesIO(rel.target_part.blob))
        else:
            ids[rel.rId]=target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,is_external=rel.is_external)
    for child in source._element:
        cloned=deepcopy(child)
        for el in cloned.iter():
            for attr,value in list(el.attrib.items()):
                if attr.startswith('{'+R+'}') and value in ids:el.set(attr,ids[value])
        target._element.append(cloned)

def _field(par,name):
    for kind in ['begin','instruction','end']:
        run=OxmlElement('w:r')
        if kind=='instruction':
            el=OxmlElement('w:instrText');el.text=' '+name+' '
        else:
            el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
        run.append(el);par._p.append(run)

def _anchor_brand(header):
    for par in header.paragraphs:
        par.paragraph_format.space_after=Pt(0)
        for run in par.runs:run.font.size=Pt(1)
    for anchor in header._element.iter('{'+WP+'}anchor'):
        anchor.set('behindDoc','1')
        for axis,align in [('positionH','left'),('positionV','top')]:
            pos=anchor.find('{'+WP+'}'+axis);pos.set('relativeFrom','page')
            for child in list(pos):pos.remove(child)
            el=OxmlElement('wp:align');el.text=align;pos.append(el)

def _control_header(section,metadata):
    width=section.page_width-section.left_margin-section.right_margin
    table=section.header.add_table(rows=3,cols=2,width=width);table.autofit=False
    rows=[['SITOP  |  '+metadata['title'],metadata['code']+'  |  v'+metadata['version']],
          ['Creado: '+metadata['created'],'Actualizado: '+metadata['updated']+'  |  BORRADOR'],
          ['Elaborado: '+metadata['prepared'],'Revisado: '+metadata['reviewed']+'  |  Aprobado: '+metadata['approved']]]
    for column in table.columns:column.width=int(width/2)
    for row,values in zip(table.rows,rows):
        for cell,value in zip(row.cells,values):
            cell.width=int(width/2);cell.text=value
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F7');cell._tc.get_or_add_tcPr().append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2)
                for run in par.runs:run.font.name='Arial';run.font.size=Pt(8)

def _page_footer(section,metadata):
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run(metadata['code']+' | v'+metadata['version']+' | Página ')
    _field(foot,'PAGE');foot.add_run(' de ');_field(foot,'NUMPAGES')
    for run in foot.runs:run.font.name='Arial';run.font.size=Pt(8)

def _section(doc,section,reference,metadata,cover=False):
    ref=Document(reference).sections[0]
    for attr in ['page_width','page_height','orientation','left_margin','right_margin']:
        setattr(section,attr,getattr(ref,attr))
    section.top_margin=Inches(1.9);section.bottom_margin=Inches(1.25)
    section.header_distance=Inches(1.05);section.footer_distance=Inches(.72)
    section.different_first_page_header_footer=cover
    for el in list(section._sectPr.findall(qn('w:pgNumType'))):section._sectPr.remove(el)
    section.header.is_linked_to_previous=False
    _clone_part_content(ref.header,section.header);_anchor_brand(section.header)
    if cover:
        # Only cover has a separate header with art and no document control.
        section.first_page_header.is_linked_to_previous=False
        _clone_part_content(section.header,section.first_page_header)
        section.first_page_footer.is_linked_to_previous=False
        for child in list(section.first_page_footer._element):section.first_page_footer._element.remove(child)
        section.first_page_footer._element.append(OxmlElement('w:p'))
    _control_header(section,metadata)
    section.footer.is_linked_to_previous=False
    for child in list(section.footer._element):section.footer._element.remove(child)
    section.footer._element.append(OxmlElement('w:p'))
    _page_footer(section,metadata)

def new_document(reference,title,code='Código pendiente',version='0.1',created=None,updated=None,
                 prepared='Pendiente',reviewed='Pendiente',approved='Pendiente'):
    """Create a cover plus a body page; cover counts as physical page 1, hidden."""
    doc=Document(reference)
    for el in list(doc._element.body):
        if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
    defaults=Document()
    if 'Table Grid' not in doc.styles:doc.styles.element.append(deepcopy(defaults.styles['Table Grid']._element))
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        if name not in doc.styles:doc.styles.element.append(deepcopy(defaults.styles[name]._element))
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
        for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Title'].font.size=Pt(24)
    for name,size in [('Heading 1',16),('Heading 2',13)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.color.rgb=RGBColor.from_string('4B438F')
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.settings.odd_and_even_pages_header_footer=False
    metadata={'title':title,'code':code,'version':version,'created':created or date.today().isoformat(),
              'updated':updated or created or date.today().isoformat(),'prepared':prepared,'reviewed':reviewed,'approved':approved}
    _section(doc,doc.sections[0],reference,metadata,True)
    titlepar=doc.add_paragraph(title,'Title');titlepar.paragraph_format.space_before=Pt(25)
    doc.add_paragraph('Oficina de software SITOP','Subtitle')
    doc.add_paragraph('BORRADOR PARA REVISIÓN  |  '+code+'  |  Versión '+version)
    table=doc.add_table(rows=0,cols=2);table.autofit=False
    width=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    for column in table.columns:column.width=int(width/2)
    for key,value in [('Creado',metadata['created']),('Actualizado',metadata['updated']),('Elaborado por',prepared),('Revisado por',reviewed),('Aprobado por',approved)]:
        cells=table.add_row().cells
        for cell,content in zip(cells,[key,value]):
            cell.width=int(width/2);cell.text=content
            for run in cell.paragraphs[0].runs:run.font.size=Pt(10)
    doc.add_paragraph('Datos institucionales y firmas pendientes de confirmación. Completar el código oficial según el proceso aplicable.')
    doc.add_page_break()
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings._element.append(update)
    doc.core_properties.title=title;doc.core_properties.author='';doc.core_properties.last_modified_by=''
    return doc,metadata

def add_orientation(doc,reference,metadata):
    """Add a body section using the other original layout, continuous PAGE."""
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    _section(doc,section,reference,metadata,False)
    return section

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reference',type=Path,required=True)
    parser.add_argument('--alternate-reference',type=Path)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--title',default='Documento de SITOP')
    parser.add_argument('--code',default='Código pendiente')
    parser.add_argument('--version',default='0.1')
    parser.add_argument('--demo',action='store_true',help='Add synthetic pages for local layout verification')
    args=parser.parse_args()
    if args.output.resolve()==args.reference.resolve() or (args.alternate_reference and args.output.resolve()==args.alternate_reference.resolve()):
        parser.error('La salida debe ser una copia; no sobrescribir referencias')
    doc,metadata=new_document(args.reference,args.title,args.code,args.version)
    if args.demo:
        for i in range(2):
            if i:doc.add_page_break()
            doc.add_heading('Contenido de verificación '+str(i+1),1)
            doc.add_paragraph('Ejemplo sintético de estructura documental. No acredita requisitos, pruebas funcionales, firmas ni aceptación institucional.')
            doc.add_heading('Control de cambios',2)
            table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
            for c,value in zip(table.rows[0].cells,['Versión','Cambio','Revisión']):c.text=value
            for c,value in zip(table.add_row().cells,[args.version,'Preparación de estructura y formato','Pendiente']):c.text=value
        if args.alternate_reference:
            add_orientation(doc,args.alternate_reference,metadata)
            doc.add_heading('Vista con orientación alternativa',1)
            doc.add_paragraph('El membrete corresponde a esta orientación y la numeración continúa. No se reutiliza un fondo deformado.')
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    print(args.output)

if __name__=='__main__':main()
